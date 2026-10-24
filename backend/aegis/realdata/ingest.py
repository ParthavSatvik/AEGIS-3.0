"""Validated import of real historical demand / generation time series (CSV, XLSX or NPP JSON).

This module is deliberately separate from the decision engine: imported series are *context data* (e.g. All-India
demand in MW), never campus telemetry and never simulator truth. Only a normalised, dimensionless shape is ever
used, by realdata/replay.py, and only to parameterise clearly labelled simulated stress scenarios.
"""
from __future__ import annotations

import csv
import datetime as dt
import hashlib
import io
import json
import re
import statistics
from dataclasses import dataclass, field

from ..settings import settings

UNITS_TO_MW = {"MW": 1.0, "GW": 1000.0, "KW": 0.001}
TS_FORMATS = ["%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M",
              "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d-%m-%Y %H:%M:%S", "%d-%m-%Y %H:%M"]
TS_HINT = re.compile(r"time|date|timestamp|updated", re.I)
VALUE_HINT = re.compile(r"demand|generation|mw|gw|value|load", re.I)


class IngestError(ValueError):
    def __init__(self, message: str, details: dict | None = None):
        super().__init__(message)
        self.details = details or {}


@dataclass
class Table:
    columns: list[str]
    rows: list[list[str]]
    fmt: str


@dataclass
class Dataset:
    id: str
    name: str
    source: str
    filename: str
    sha256: str
    unit_in: str
    columns: dict
    points: list[tuple[dt.datetime, float]] = field(repr=False)
    quality: dict = field(default_factory=dict)

    def meta(self) -> dict:
        return {"id": self.id, "name": self.name, "source": self.source, "filename": self.filename,
                "sha256": self.sha256, "unit_in": self.unit_in, "unit": "MW", "columns": self.columns,
                "quality": self.quality,
                "kind": "real historical context data (not campus telemetry, not simulator truth)"}


def read_table(filename: str, data: bytes) -> Table:
    """Parse CSV / XLSX / NPP-JSON into string cells, enforcing size and row limits."""
    if len(data) > settings.max_upload_bytes:
        raise IngestError(f"file too large: {len(data)} bytes > limit {settings.max_upload_bytes}")
    if not data:
        raise IngestError("empty file")
    name = filename.lower()
    if name.endswith(".json"):
        try:
            obj = json.loads(data)
        except json.JSONDecodeError as e:
            raise IngestError(f"invalid JSON: {e.msg}") from None
        if not (isinstance(obj, list) and obj and all(isinstance(x, dict) for x in obj)):
            raise IngestError("JSON must be a non-empty list of records")
        cols = list(obj[0].keys())
        rows = [[_json_cell(x.get(c)) for c in cols] for x in obj]
        table = Table(cols, rows, "json")
    elif name.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            ws = wb.worksheets[0]
            it = ws.iter_rows(values_only=True)
            header = next(it)
            rows = []
            for r in it:
                rows.append(["" if v is None else (v.strftime("%Y-%m-%d %H:%M:%S") if isinstance(v, dt.datetime)
                                                   else str(v)) for v in r])
                if len(rows) > settings.max_upload_rows:
                    break
        except IngestError:
            raise
        except Exception as e:  # corrupted or unsupported workbook
            raise IngestError(f"could not read XLSX: {type(e).__name__}") from None
        table = Table([str(h) if h is not None else f"column_{i + 1}" for i, h in enumerate(header)], rows, "xlsx")
    elif name.endswith((".csv", ".txt")):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise IngestError("CSV must be UTF-8 encoded") from None
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(text), dialect)
        all_rows = [r for r in reader if any(c.strip() for c in r)]
        if not all_rows:
            raise IngestError("no rows found")
        table = Table([c.strip() for c in all_rows[0]], all_rows[1:], "csv")
    else:
        raise IngestError("unsupported file type: use .csv, .xlsx or .json")
    if len(table.rows) > settings.max_upload_rows:
        raise IngestError(f"too many rows: more than {settings.max_upload_rows}")
    if not table.rows:
        raise IngestError("file has a header but no data rows")
    return table


def _json_cell(v) -> str:
    if isinstance(v, (int, float)) and not isinstance(v, bool) and v > 1e11:  # epoch milliseconds (NPP format)
        return dt.datetime.fromtimestamp(v / 1000, dt.timezone(dt.timedelta(hours=5, minutes=30))) \
            .strftime("%Y-%m-%d %H:%M:%S")
    return "" if v is None else str(v)


def suggest_mapping(table: Table) -> dict:
    ts = next((c for c in table.columns if TS_HINT.search(c)), None)
    val = next((c for c in table.columns if c != ts and VALUE_HINT.search(c)), None)
    unit = None
    if val:
        m = re.search(r"\b(MW|GW|kW)\b", val, re.I)
        unit = m.group(1).upper() if m else None
    return {"timestamp_column": ts, "value_column": val, "unit": unit}


def preview(filename: str, data: bytes) -> dict:
    t = read_table(filename, data)
    return {"format": t.fmt, "columns": t.columns, "row_count": len(t.rows), "first_rows": t.rows[:8],
            "suggested": suggest_mapping(t)}


def _parse_ts(values: list[str]) -> tuple[list[dt.datetime | None], str | None]:
    """Pick the single format that parses the most rows; one format for the whole file (no guessing per row)."""
    best: tuple[int, str | None, list] = (0, None, [])
    for f in TS_FORMATS:
        out = []
        for v in values:
            try:
                out.append(dt.datetime.strptime(v.strip(), f))
            except ValueError:
                out.append(None)
        n = sum(x is not None for x in out)
        if n > best[0]:
            best = (n, f, out)
    return best[2] or [None] * len(values), best[1]


def build_dataset(ds_id: str, filename: str, data: bytes, source: str, timestamp_column: str,
                  value_column: str, unit: str | None, name: str = "") -> Dataset:
    t = read_table(filename, data)
    missing = [c for c in (timestamp_column, value_column) if c not in t.columns]
    if missing:
        raise IngestError(f"column(s) not found: {missing}", {"columns": t.columns})
    if timestamp_column == value_column:
        raise IngestError("timestamp and value columns must differ")
    unit = (unit or suggest_mapping(Table([value_column], [], "")).get("unit") or "").upper()
    if unit not in UNITS_TO_MW:
        raise IngestError("unit required: choose MW, GW or kW (could not be read from the column header)")
    if not source.strip():
        raise IngestError("source attribution is required")
    ti, vi = t.columns.index(timestamp_column), t.columns.index(value_column)
    raw_ts = [r[ti] if ti < len(r) else "" for r in t.rows]
    raw_v = [r[vi] if vi < len(r) else "" for r in t.rows]
    stamps, fmt = _parse_ts(raw_ts)
    bad_ts = sum(1 for s in stamps if s is None)
    if fmt is None or bad_ts > 0.5 * len(stamps):
        raise IngestError("timestamps could not be parsed", {"examples": raw_ts[:3], "accepted_formats": TS_FORMATS})
    missing_v, invalid_v, negative = 0, 0, 0
    pts: dict[dt.datetime, float] = {}
    dupes = 0
    for s, v in zip(stamps, raw_v, strict=True):
        if s is None:
            continue
        txt = v.strip().replace(",", "")
        if txt == "" or txt.lower() in ("na", "nan", "null", "-"):
            missing_v += 1
            continue
        try:
            x = float(txt)
        except ValueError:
            invalid_v += 1
            continue
        if x < 0:
            negative += 1
            continue
        if s in pts:
            dupes += 1
            continue
        pts[s] = x * UNITS_TO_MW[unit]
    series = sorted(pts.items())
    if len(series) < 2:
        raise IngestError("fewer than 2 valid data points after validation")
    gaps = [(b[0] - a[0]).total_seconds() / 60 for a, b in zip(series, series[1:], strict=False)]
    med = statistics.median(gaps)
    vals = [v for _, v in series]
    quality = {
        "rows_in_file": len(t.rows), "valid_points": len(series), "duplicate_timestamps_dropped": dupes,
        "unparseable_timestamps": bad_ts, "missing_values": missing_v, "non_numeric_values": invalid_v,
        "negative_values_rejected": negative, "timestamp_format": fmt,
        "timezone": "as recorded in the file (NPP JSON converted to IST, UTC+05:30)",
        "start": series[0][0].isoformat(), "end": series[-1][0].isoformat(),
        "median_interval_min": round(med, 2),
        "gaps_over_3x_median": sum(1 for g in gaps if g > 3 * med),
        "min_mw": round(min(vals), 3), "max_mw": round(max(vals), 3), "mean_mw": round(statistics.fmean(vals), 3),
    }
    return Dataset(id=ds_id, name=name or filename, source=source.strip()[:300], filename=filename[:200],
                   sha256=hashlib.sha256(data).hexdigest(), unit_in=unit,
                   columns={"timestamp": timestamp_column, "value": value_column}, points=series, quality=quality)


def trend(ds: Dataset, max_points: int = 500) -> list[dict]:
    step = max(1, len(ds.points) // max_points)
    return [{"t": t.isoformat(), "mw": round(v, 3)} for t, v in ds.points[::step]]
